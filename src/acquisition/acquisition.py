"""
=============================================================================
ระบบควบคุมการดึงข้อมูลและจัดเก็บไฟล์ดิบ (Data Acquisition Engine - Step 1)
=============================================================================
หลักการทำงาน (Architectural Highlights):
1. **Approved Sources Only (Strict Rule)**:
   - ตรวจสอบคอลัมน์ status ใน Catalog โดยประมวลผล "เฉพาะชุดข้อมูลที่ได้รับอนุมัติ (approved)" เท่านั้น
   - ชุดข้อมูลที่มีสถานะ candidate หรือ rejected จะถูกปฏิเสธ/ข้ามโดยอัตโนมัติ เพื่อรักษาคุณภาพของข้อมูล
2. **Modular Strategy Pattern (Clean Architecture)**:
   - ตรรกะการเชื่อมต่อเว็บไซต์แต่ละประเภทถูกแยกเป็นโมดูลย่อยใน `extractors/`:
     - portal_extractor.py (Open Data Portals: data.go.th, eppo.go.th, catalog.dmf.go.th, doeb.go.th)
     - direct_extractor.py (ลิงก์ตรง .csv, .xlsx, .pdf, .zip)
     - form_extractor.py (เว็บแอปพลิเคชันแบบฟอร์ม: thaicarbonlabel.tgo.or.th และ dmf.go.th/public/supply)
3. **Idempotency & Cryptographic Audit**:
   - คำนวณ SHA-256 ทุกไฟล์ และลงบันทึกใน `data/raw/_provenance_ledger.jsonl`
   - หากมีไฟล์เดิมอยู่แล้วและ Checksum ถูกต้อง จะข้ามการดาวน์โหลดซ้ำเพื่อประหยัดแบนด์วิดท์
=============================================================================
"""

import csv
import os
import sys
from pathlib import Path
from urllib.parse import urlparse

# กำหนด Console encoding สำหรับ Windows ให้แสดงผลภาษาไทยได้ถูกต้อง
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

CURRENT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = CURRENT_DIR.parent.parent
RAW_DATA_DIR = PROJECT_ROOT / "data" / "raw"
EXTRACTORS_DIR = CURRENT_DIR / "extractors"

# นำเข้าโมดูลย่อยโดยไม่ต้องพึ่ง __init__.py
if str(CURRENT_DIR) not in sys.path:
    sys.path.append(str(CURRENT_DIR))
if str(EXTRACTORS_DIR) not in sys.path:
    sys.path.append(str(EXTRACTORS_DIR))

from provenance import (
    record_provenance,
    load_raw_metadata,
    compute_sha256,
    generate_batch_id,
    verify_raw_integrity,
)
from base_extractor import BaseExtractor, route_and_extract


def get_catalog_path() -> Path:
    """
    ค้นหาไฟล์ทะเบียนแหล่งข้อมูลหลัก (Source Registry)
    ใช้ไฟล์ Excel (.xlsx) ที่ docs/source_registry/source_registry.xlsx เป็น Single Source of Truth
    """
    candidates = [
        PROJECT_ROOT / "docs" / "source_registry" / "source_registry.xlsx",
        PROJECT_ROOT / "docs" / "Source Registry" / "source_registry.xlsx",
        PROJECT_ROOT / "source_registry.xlsx",
    ]
    for p in candidates:
        if p.exists():
            return p
    raise FileNotFoundError(
        "ไม่พบไฟล์ Source Registry (docs/source_registry/source_registry.xlsx) ในโปรเจกต์"
    )


def load_all_catalog_records() -> list[dict]:
    """
    โหลดข้อมูลทุกแถวใน Catalog จากไฟล์ Excel (docs/source_registry/source_registry.xlsx)
    โดยอ่านชีต 'Source Registry' โดยตรง
    """
    path = get_catalog_path()

    # โหลดไฟล์ Excel (.xlsx / .xls)
    if path.suffix.lower() in [".xlsx", ".xls"]:
        try:
            import openpyxl
        except ImportError:
            raise ImportError(
                "\n" + "=" * 60 + "\n"
                "[ERROR] ไม่พบไลบรารี 'openpyxl' ใน Python ตัวที่กำลังรันอยู่!\n\n"
                "[FIX] วิธีแก้ไขทันที (พิมพ์คำสั่งนี้ใน Terminal):\n"
                "   python -m pip install openpyxl\n\n"
                "[FIX] หรือติดตั้ง Dependencies ทั้งหมด:\n"
                "   python -m pip install -r requirements.txt\n" + "=" * 60
            )
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        target_sheet = "Source Registry" if "Source Registry" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[target_sheet]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return []

        header = [str(col).strip() if col is not None else "" for col in rows[0]]
        records = []
        for r in rows[1:]:
            if not any(r):
                continue
            d = {
                header[i]: (str(val).strip() if val is not None else "")
                for i, val in enumerate(r)
                if i < len(header) and header[i]
            }
            if d.get("source_id"):
                records.append(d)

        return records

    # กรณีฉุกเฉิน Fallback สำหรับไฟล์ CSV
    with open(path, mode="r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        return [dict(r) for r in reader]


def get_approved_sources() -> list[str]:
    """ดึงรายชื่อรหัส source_id ทั้งหมดที่มีสถานะ 'approved' จาก Catalog"""
    records = load_all_catalog_records()
    return [
        r["source_id"].strip()
        for r in records
        if r.get("status", "").strip().lower() == "approved"
    ]


def load_source_from_catalog(target_source_id: str) -> dict:
    """อ่านข้อมูลของ source_id ที่ระบุจาก Catalog"""
    records = load_all_catalog_records()
    for row in records:
        if row.get("source_id", "").strip().upper() == target_source_id.strip().upper():
            return row
    catalog_path = get_catalog_path()
    raise ValueError(f"ไม่พบรหัส {target_source_id} ใน {catalog_path.name}")


# ตารางจับคู่ชื่อหน่วยงานไทย -> รหัส Owner Code สากล
AGENCY_MAP = {
    "กรมเชื้อเพลิงธรรมชาติ": "DMF",
    "กรมธุรกิจพลังงาน": "DOEB",
    "สำนักงานนโยบายและแผนพลังงาน": "EPPO",
    "กรมพัฒนาพลังงานทดแทนและอนุรักษ์พลังงาน": "DEDE",
    "สำนักงานปลัดกระทรวงพลังงาน": "OPS",
    "สำนักงานสถิติแห่งชาติ": "NSO",
    "ธนาคารแห่งประเทศไทย": "BOT",
    "กรมสรรพากร": "RD",
    "กรมสรรพสามิต": "EXCISE",
    "องค์การบริหารจัดการก๊าซเรือนกระจก": "TGO",
}


def resolve_owner_code(publisher: str, url: str) -> str:
    """วิเคราะห์และกำหนดรหัสย่อของหน่วยงาน (Owner Code) อย่างแม่นยำ"""
    # 1. ตรวจสอบจากชื่อ Publisher ใน Catalog
    for th_name, code in AGENCY_MAP.items():
        if th_name in publisher:
            return code

    # 2. กรณี Publisher ว่าง ให้ตรวจสอบจาก Domain Name ของ URL
    domain = urlparse(url).netloc.lower()
    if "tgo.or.th" in domain:
        return "TGO"
    elif "eppo.go.th" in domain:
        return "EPPO"
    elif "dmf.go.th" in domain:
        return "DMF"
    elif "doeb.go.th" in domain:
        return "DOEB"
    elif "bot.or.th" in domain:
        return "BOT"

    return "DMF"


def detect_flow_from_title(title: str) -> str:
    """
    วิเคราะห์ชื่อชุดข้อมูลเพื่อสร้างรหัส Flow หมวดหมู่ข้อมูลตามมาตรฐาน
    (ตรวจสอบหมวดหมู่ Carbon ก่อนเพื่อไม่ให้ 'ผลิตภัณฑ์' ไปตรงกับ 'ผลิต')
    """
    t = title.lower()
    if "คาร์บอน" in t or "co2" in t or "carbon" in t:
        return "Carbon-Footprint"
    elif "ขาย" in t or "จำหน่าย" in t:
        return "Petroleum-Sales"
    elif "ใช้" in t or "บริโภค" in t:
        return "Petroleum-Consumption"
    elif "จัดหา" in t:
        return "Petroleum-Supply"
    elif "สำรอง" in t:
        return "Petroleum-Reserve"
    elif "นำเข้า" in t:
        return "Petroleum-Import"
    elif "ขนส่ง" in t or "ขนถ่าย" in t:
        return "Petroleum-Transport"
    elif "ราคา" in t:
        return "Petroleum-Price"
    elif "ผลิต" in t:
        return "Petroleum-Production"
    elif "โครงการ" in t:
        return "Off-Budget-Project"
    return "Petroleum-Report"


def generate_raw_filename(
    source_id: str,
    owner: str,
    flow: str,
    year_be: int,
    month: int,
    extension: str = "csv",
) -> str:
    """
    สร้างชื่อไฟล์ดิบตามข้อกำหนด:
    {source_id}_{owner}_{flow}_TH_{year_ce}{month:02d}_RAW.{extension}
    ตัวอย่าง: SRC-033_EPPO_Petroleum-Consumption_TH_202408_RAW.csv
    """
    year_ce = year_be - 543
    clean_ext = extension.replace(".", "").lower()
    return f"{source_id}_{owner}_{flow}_TH_{year_ce}{month:02d}_RAW.{clean_ext}"


def _is_approved_gate(source_info: dict) -> bool:
    """[STRICT RULE] ตรวจสอบว่าชุดข้อมูลมีสถานะ approved หรือไม่"""
    status = source_info.get("status", "").strip().lower()
    if status != "approved":
        sid = source_info.get("source_id", "UNKNOWN")
        print(f"  [SKIPPED] ข้ามชุดข้อมูล {sid}:")
        print(f"   สถานะปัจจุบันคือ '{status}' (ระบบเปิดรับเฉพาะชุดข้อมูลที่ได้รับอนุมัติ 'approved' เท่านั้น)")
        if source_info.get("notes"):
            print(f"   หมายเหตุ: {source_info.get('notes')}")

        return False
    return True


def _check_idempotency_cache(file_path: Path, force: bool = False) -> dict | None:
    """ตรวจสอบว่ามีไฟล์ดิบเดิมอยู่แล้วและ Checksum ถูกต้องหรือไม่ (ป้องกันดาวน์โหลดซ้ำ)"""
    if file_path.exists() and not force:
        print(f"  [CHECK] พบไฟล์ดิบเดิมอยู่แล้ว: {file_path.name}")
        integrity = verify_raw_integrity(str(file_path))
        if integrity.get("is_valid"):
            print(f"  [SKIP] ตรวจสอบ SHA-256 ผ่านสมบูรณ์แล้ว (ข้ามดาวน์โหลดซ้ำ)")
            
            try:
                return load_raw_metadata(str(file_path))
            except Exception:
                return {"status": "cached", "file_path": str(file_path)}
    return None


def _save_raw_file(output_path: Path, content_bytes: bytes) -> float:
    """บันทึกไฟล์ไบนารีสภาพดิบ 100% ลงใน data/raw/"""
    with open(output_path, "wb") as f:
        f.write(content_bytes)
    file_size_kb = len(content_bytes) / 1024
    print(f"  [OK] บันทึกไฟล์ดิบสำเร็จ: {output_path.name} ({file_size_kb:.2f} KB)")
    return file_size_kb


# =============================================================================
# ฟังก์ชันหลัก: ดึงข้อมูลดิบ (Data Acquisition Engine)
# =============================================================================
def acquire_source(
    source_id: str = "SRC-033",
    year_be: int = 2567,
    month: int = 8,
    flow: str = None,
    owner: str = None,
    force: bool = False,
) -> dict:
    """
    ดึงข้อมูลและจัดเก็บไฟล์ดิบตามมาตรฐาน Step 1 (Data Acquisition)
    - **กฎสำคัญ:** ดึงเฉพาะชุดข้อมูลที่มีสถานะ 'approved' เท่านั้น
    - รองรับสถาปัตยกรรมเว็บไซต์ทุกรูปแบบผ่าน Modular Extractors
    - มีระบบ Idempotency ป้องกันดาวน์โหลดซ้ำ
    - บันทึก Checksum SHA-256 ลงใน _provenance_ledger.jsonl
    """




    # 1. โหลดข้อมูลจาก Catalog
    source_info = load_source_from_catalog(source_id)

    # 2. [STRICT GATE] ตรวจสอบสถานะการอนุมัติ
    if not _is_approved_gate(source_info):
        status = source_info.get("status", "").strip().lower()
        return {
            "status": "skipped",
            "source_id": source_id,
            "reason": f"Status '{status}' is not approved",
        }

    # 3. วิเคราะห์ Metadata
    title = source_info.get("title", "")
    page_url = source_info.get("url_or_document_id", "")
    expected_type = source_info.get("source_type", "CSV").strip().upper()
    publisher = source_info.get("publisher_or_owner", "")
    resolved_owner = owner or resolve_owner_code(publisher, page_url)
    resolved_flow = flow or detect_flow_from_title(title)

    print(f"  [INFO] {title} | {resolved_owner} ({resolved_flow}) | {expected_type}")

    RAW_DATA_DIR.mkdir(parents=True, exist_ok=True)

    # 4. กรณีพิเศษสำหรับ TGO (SRC-043): บันทึก 2 ตารางแยกอิสระ (ปิโตรเลียม & ปิโตรเคมี)
    if source_id == "SRC-043":
        fn_petro = RAW_DATA_DIR / f"{source_id}_{resolved_owner}_Carbon-Footprint-Petroleum_TH_{year_be - 543}{month:02d}_RAW.html"
        fn_chem = RAW_DATA_DIR / f"{source_id}_{resolved_owner}_Carbon-Footprint-Petrochemical_TH_{year_be - 543}{month:02d}_RAW.html"
        
        if fn_petro.exists() and fn_chem.exists() and not force:
            chk1 = verify_raw_integrity(str(fn_petro))
            chk2 = verify_raw_integrity(str(fn_chem))
            if chk1.get("is_valid") and chk2.get("is_valid"):
                print(f"  [CHECK] พบไฟล์ดิบ TGO ทั้ง 2 ตารางเดิมอยู่แล้ว:")
                print(f"    1. {fn_petro.name}")
                print(f"    2. {fn_chem.name}")
                print(f"  [SKIP] ตรวจสอบ SHA-256 ผ่านสมบูรณ์ทั้ง 2 ตาราง (ข้ามดาวน์โหลดซ้ำ)")
                return load_raw_metadata(str(fn_petro))

        from extractors.form_extractor import FormExtractor
        tgo_engine = FormExtractor()
        tables = tgo_engine.extract_tgo_tables()
        if "Petroleum" in tables and "Petrochemical" in tables:
            c_petro, u_petro = tables["Petroleum"]
            c_chem, u_chem = tables["Petrochemical"]

            _save_raw_file(fn_petro, c_petro)
            record_provenance(
                source_id=source_id,
                raw_file_path=str(fn_petro),
                source_url=u_petro,
                page_or_table="คาร์บอนฟุตพริ้นท์ผลิตภัณฑ์ (ปิโตรเลี่ยม)",
                extraction_method="tgo_targeted_form_extractor",
            )

            _save_raw_file(fn_chem, c_chem)
            receipt = record_provenance(
                source_id=source_id,
                raw_file_path=str(fn_chem),
                source_url=u_chem,
                page_or_table="คาร์บอนฟุตพริ้นท์ผลิตภัณฑ์ (ปิโตรเคมี และเคมีภัณฑ์)",
                extraction_method="tgo_targeted_form_extractor",
            )
            
            return receipt

    # 4. กำหนดชื่อไฟล์ที่คาดหวัง และตรวจสอบ Idempotency Cache (ป้องกันดาวน์โหลดซ้ำ)
    # สำหรับชุดข้อมูล Master ย้อนหลังระยะยาวจาก Open Data Catalog (ครอบคลุม พ.ศ. 2546/01 ถึง 2569/03 = ค.ศ. 200301-202603)
    if source_id in ["SRC-036", "SRC-044", "SRC-045", "SRC-046"]:
        period_str = "200301-202603"
        expected_filename = f"{source_id}_{resolved_owner}_{resolved_flow}_TH_{period_str}_RAW.{expected_type.lower()}"
    else:
        expected_filename = generate_raw_filename(
            source_id=source_id,
            owner=resolved_owner,
            flow=resolved_flow,
            year_be=year_be,
            month=month,
            extension=expected_type,
        )

    cached_receipt = _check_idempotency_cache(RAW_DATA_DIR / expected_filename, force=force)
    if cached_receipt:
        return cached_receipt

    # 5. สั่งงาน Extractor ผ่าน Universal Router
    content_bytes, direct_url, actual_ext = route_and_extract(
        url=page_url,
        preferred_type=expected_type,
        year_be=year_be,
        month=month,
    )

    if not content_bytes or len(content_bytes) == 0:
        raise ValueError(f"ไม่สามารถดาวน์โหลดเนื้อหาจาก: {page_url}")

    # 6. บันทึกไฟล์ดิบ (100% Untouched Raw Data)
    if source_id in ["SRC-036", "SRC-044", "SRC-045", "SRC-046"]:
        period_str = "200301-202603"
        raw_filename = f"{source_id}_{resolved_owner}_{resolved_flow}_TH_{period_str}_RAW.{actual_ext.lower()}"
    else:
        raw_filename = generate_raw_filename(
            source_id=source_id,
            owner=resolved_owner,
            flow=resolved_flow,
            year_be=year_be,
            month=month,
            extension=actual_ext,
        )
    output_path = RAW_DATA_DIR / raw_filename
    _save_raw_file(output_path, content_bytes)

    # 7. บันทึก Checksum ลงสมุดคุมประวัติ (_provenance_ledger.jsonl)
    receipt = record_provenance(
        source_id=source_id,
        raw_file_path=str(output_path),
        source_url=direct_url,
        page_or_table=title,
        extraction_method="modular_extractor_engine",
    )

    
    return receipt


# =============================================================================
# ฟังก์ชันสั่งรันชุดข้อมูล Approved ทั้งหมด หรือผ่านคำสั่ง CLI
# =============================================================================
def acquire_all_approved(force: bool = False):
    """สั่งดาวน์โหลดชุดข้อมูลที่ได้รับอนุมัติ (Approved Sources) ทั้งหมดใน Catalog"""
    approved = get_approved_sources()
    print(f"\n========================================================")
    print(f"[START] เริ่มต้นดาวน์โหลดชุดข้อมูลที่ผ่านการอนุมัติ (Approved) ทั้งหมด {len(approved)} รายการ:")
    print(f"        {', '.join(approved)}")
    print(f"========================================================")
    
    success_count = 0
    for sid in approved:
        try:
            res = acquire_source(source_id=sid, force=force)
            if res and res.get("status") != "skipped":
                success_count += 1
        except Exception as e:
            print(f"[ERROR] ไม่สามารถดึง {sid}: {e}\n")

    print(f"[SUMMARY] สรุปผล: ดาวน์โหลดชุดข้อมูล Approved สำเร็จ {success_count}/{len(approved)} รายการ\n")


def main():
    args = sys.argv[1:]
    force = "--force" in args
    clean_args = [a for a in args if not a.startswith("--")]

    # กรณีสั่งดูรายการ approved: python src/acquisition/acquisition.py --list
    if "--list" in args:
        approved = get_approved_sources()
        print(f"\n[LIST] รายการ Approved Sources ({len(approved)} รายการ):")
        for sid in approved:
            info = load_source_from_catalog(sid)
            print(f"  - {sid}: {info.get('title')} [{info.get('publisher_or_owner')}]")
        print()
        return

    # หากไม่มีการระบุ source_id มาใน Terminal ให้รันชุด Approved ทั้งหมดโดยอัตโนมัติ
    if not clean_args:
        acquire_all_approved(force=force)
        return

    # หากมีการระบุ source_id เจาะจง ให้ประมวลผลตามที่สั่ง (มีระบบดักสถานะ approved ให้)
    print(f"\n[BATCH] รับคำสั่งดาวน์โหลด: {', '.join(clean_args)}")
    for sid in clean_args:
        try:
            acquire_source(source_id=sid, force=force)
        except Exception as e:
            print(f"[ERROR] เกิดข้อผิดพลาดกับ {sid}: {e}\n")


if __name__ == "__main__":
    main()
